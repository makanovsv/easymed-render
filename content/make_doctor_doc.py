"""Word-файл врачу: памятки (ru + kk рядом). python3 make_doctor_doc.py <выход.docx> <id,id,...|@файл со строками каталога> [заголовок]"""
import json, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT

LBL = [('what', 'Что это', 'Бұл не'), ('why', 'Зачем', 'Не үшін'), ('duration', 'Сколько длится', 'Қанша уақыт'),
       ('tell_doctor', 'Скажите врачу', 'Дәрігерге айтыңыз'), ('prepare', 'Подготовка', 'Дайындық'),
       ('bring', 'Что взять с собой', 'Өзіңізбен не алу керек'), ('steps', 'Как проходит', 'Қалай өтеді'),
       ('feel', 'Что вы почувствуете', 'Не сезінесіз'), ('after', 'После', 'Кейін'),
       ('red_flags', 'Срочно к врачу / 103', 'Шұғыл дәрігерге / 103'), ('faq', 'Вопросы и ответы', 'Сұрақтар мен жауаптар'),
       ('reminders', 'Напоминание', 'Еске салу')]


def fmt(v):
    if isinstance(v, str):
        return v
    out = []
    for x in v:
        if isinstance(x, str): out.append('• ' + x)
        elif 'q' in x: out.append(f"— {x['q']}\n  {x['a']}")
        elif 'title' in x: out.append(f"• {x['title']}: {x['text']}")
        elif 'text' in x: out.append(f"• за {x.get('hoursBefore', '')} ч: {x['text']}")
    return '\n'.join(out)


def main(out, ids, title):
    C, K = Path('cards'), Path('cards_kk')
    cards = [json.loads((C / f'{i}.json').read_text()) for i in ids]
    order = {'H': 0, 'M': 1, 'L': 2}
    cards.sort(key=lambda c: (order[c['review']], c['section'], c['title']))
    d = Document()
    s = d.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = s.page_height, s.page_width
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(s, m, Cm(1.6))
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    d.add_heading(title, 0)
    nh = sum(c['review'] == 'H' for c in cards)
    d.add_paragraph('Памятки для пациентов: что сказать врачу, как подготовиться, что будет в кабинете и что после. По ним делаются видео на 9 языках и страницы сайта. Тексты простым языком, они не заменяют консультацию.')
    d.add_paragraph('Что нужно: проверить, что нет ошибок и опасных советов. Править прямо в тексте (режим «Исправления» в Word) или комментарием. Где зависит от клиники — достаточно пометить «уточнять в клинике». Казахская колонка — перевод русской.')
    d.add_paragraph(f'Всего {len(cards)}: сначала {nh} операций, родов и наркоза (без вашей подписи не публикуем), потом {len(cards) - nh} процедур.')
    for n, c in enumerate(cards, 1):
        k = {**c, **(json.loads((K / f"{c['id']}.json").read_text()) if (K / f"{c['id']}.json").exists() else {})}
        d.add_heading(f"{n}. {c['title']}  /  {k.get('title', '')}", 2)
        d.add_paragraph(f"код: {c['id']} · уровень {c['review']}").runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        t = d.add_table(rows=1, cols=3); t.style = 'Table Grid'
        for i, h in enumerate(['', 'Русский', 'Қазақша']): t.rows[0].cells[i].text = h
        for key, lr, lk in LBL:
            if not c.get(key): continue
            r = t.add_row().cells; r[0].text = f"{lr}\n{lk}"; r[1].text = fmt(c[key]); r[2].text = fmt(k.get(key, ''))
        for row in t.rows:
            row.cells[0].width = Cm(3.5); row.cells[1].width = Cm(11.8); row.cells[2].width = Cm(11.8)
            for p in row.cells[0].paragraphs:
                for r_ in p.runs: r_.bold = True
        if n < len(cards): d.add_page_break()
    d.save(out); print(out, len(cards))


if __name__ == '__main__':
    a = sys.argv[2]
    ids = [l.split('|')[0] for l in open(a[1:]) if l.strip() and not l.startswith('#') and l.split('|')[5] in 'MH'] if a.startswith('@') else a.split(',')
    main(sys.argv[1], ids, sys.argv[3] if len(sys.argv) > 3 else 'EasyMed Global — тексты на проверку врачу')
